#!/usr/bin/env python3
"""Build incremental SAT vocabulary batch JSON files from the source workbook."""
from __future__ import annotations

import argparse, json, re, unicodedata
from collections import defaultdict
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "source" / "Barrons_1100_by_root_prefix.xlsx"
DATA_DIR, BATCH_DIR = ROOT / "data", ROOT / "data" / "batches"
INDEX_PATH, PROGRESS_PATH = DATA_DIR / "batches-index.json", DATA_DIR / "progress.json"

def clean(value): return "" if value is None else str(value).strip()
def split_lines(value): return [x.strip() for x in clean(value).replace("\r\n", "\n").split("\n") if x.strip()]
def normalize_pos(value):
    value = clean(value).lower().replace(" ", "")
    if value.startswith("adj"): return "adj."
    if value.startswith("adv"): return "adv."
    if value.startswith("n"): return "n."
    if value.startswith("v"): return "v."
    return value or "other"
def slugify(value):
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-zA-Z0-9]+", "-", value).strip("-").lower() or "batch"
def sentence_pairs(en_value, zh_value):
    en, zh = split_lines(en_value), split_lines(zh_value)
    en, zh = (en + [""] * 2)[:2], (zh + [""] * 2)[:2]
    return [{"en": en[i], "zh": zh[i]} for i in range(2)]

def parse_categorized(ws, part):
    records, current = [], ["", "", ""]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(clean(row[i]) for i in range(3)):
            current = [clean(row[i]) or current[i] for i in range(3)]
        if not clean(row[3]): continue
        records.append({"word":clean(row[3]),"pos":clean(row[4]),"pos_key":normalize_pos(row[4]),
            "defn_en":clean(row[5]),"defn_zh":clean(row[6]),"root_note":clean(row[10]),
            "sentences":sentence_pairs(row[8],row[9]),
            "category":{"part":part,"num":current[0],"spelling":current[1],"meaning":current[2]}})
    return records

def parse_unclassified(ws):
    records=[]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not clean(row[0]): continue
        records.append({"word":clean(row[0]),"pos":clean(row[1]),"pos_key":normalize_pos(row[1]),
            "defn_en":clean(row[2]),"defn_zh":clean(row[3]),"root_note":clean(row[7]),
            "sentences":sentence_pairs(row[5],row[6]),"category":None})
    return records

def load_records(source):
    wb=openpyxl.load_workbook(source,data_only=True)
    p1_name=next((n for n in wb.sheetnames if n.startswith("Part 1 ")),None)
    p2_name=next((n for n in wb.sheetnames if n.startswith("Part 2 ")),None)
    u_name=next((n for n in wb.sheetnames if "Unclassified" in n),None)
    if not all((p1_name,p2_name,u_name)): raise ValueError("Expected Part 1, Part 2, and Unclassified sheets")
    return parse_categorized(wb[p1_name],1),parse_categorized(wb[p2_name],2),parse_unclassified(wb[u_name])

def categorized_batches(records,part):
    grouped={}
    for r in records:
        c=r["category"]; grouped.setdefault((c["num"],c["spelling"],c["meaning"]),[]).append(r)
    output=[]
    for (num,spelling,meaning),words in grouped.items():
        ordered=sorted(words,key=lambda x:x["word"].casefold()) if len(words)>50 else words
        chunks=[ordered[i:i+50] for i in range(0,len(ordered),50)]
        for i,chunk in enumerate(chunks,1):
            suffix=f" ({i}/{len(chunks)})" if len(chunks)>1 else ""
            batch_id=f"p{part}-{num}-{slugify(spelling)}"+(f"-{i}" if len(chunks)>1 else "")
            output.append({"id":batch_id,"part":f"part{part}","part_label":"Part 1 字首篇 (Prefixes)" if part==1 else "Part 2 字根篇 (Roots)","name":f"{spelling}{suffix}","subtitle":meaning,"count":len(chunk),"words":chunk})
    return output

def unclassified_batches(records):
    ordered=sorted(records,key=lambda x:x["word"].casefold()); output=[]
    for i in range(0,len(ordered),50):
        chunk=ordered[i:i+50]; n=i//50+1; first,last=chunk[0]["word"],chunk[-1]["word"]
        output.append({"id":f"u-{n:02d}-{slugify(first)}-{slugify(last)}","part":"unclassified","part_label":"未歸類 Unclassified","name":f"Unclassified {n:02d} ({first}–{last})","subtitle":"Alphabetical batch","count":len(chunk),"words":chunk})
    return output

def fallbacks(records,batch_words):
    current={r["word"].casefold() for r in batch_words}; pools=defaultdict(list); seen=defaultdict(set)
    for r in records:
        key,word=r["pos_key"],r["word"].casefold()
        if word in current or word in seen[key] or len(pools[key])>=12: continue
        seen[key].add(word); pools[key].append({k:r[k] for k in ("word","pos","pos_key","defn_en","defn_zh")})
    return dict(pools)

def write_json(path,value):
    path.parent.mkdir(parents=True,exist_ok=True); path.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding="utf-8")

def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--source",type=Path,default=DEFAULT_SOURCE); parser.add_argument("--batch-count",type=int,default=1); parser.add_argument("--batch-id")
    args=parser.parse_args(); p1,p2,u=load_records(args.source); all_records=p1+p2+u
    batches=categorized_batches(p1,1)+categorized_batches(p2,2)+unclassified_batches(u)
    categories=len({(r["category"]["part"],r["category"]["num"],r["category"]["spelling"]) for r in p1+p2})
    if (len(all_records),len(u),categories,len(batches))!=(6805,3824,134,236): raise ValueError(f"Sanity check failed: {len(all_records)=}, {len(u)=}, {categories=}, {len(batches)=}")
    progress={"generated":[],"total":len(batches)}
    if PROGRESS_PATH.exists(): progress=json.loads(PROGRESS_PATH.read_text(encoding="utf-8"))
    generated=set(progress.get("generated",[]))
    selected=[b for b in batches if b["id"]==args.batch_id] if args.batch_id else [b for b in batches if b["id"] not in generated][:max(0,args.batch_count)]
    if args.batch_id and not selected: raise ValueError(f"Unknown batch id: {args.batch_id}")
    for b in selected:
        payload=dict(b); payload["fallback_options"]=fallbacks(all_records,b["words"]); write_json(BATCH_DIR/f"{b['id']}.json",payload); generated.add(b["id"]); print(f"Generated {b['id']}: {b['name']} ({b['count']} words)")
    ordered=[b["id"] for b in batches if b["id"] in generated]; next_id=next((b["id"] for b in batches if b["id"] not in generated),None)
    write_json(PROGRESS_PATH,{"generated":ordered,"generated_count":len(ordered),"total":len(batches),"next_batch":next_id})
    metadata=[]
    for b in batches:
        metadata.append({"id":b["id"],"part":b["part"],"part_label":b["part_label"],"name":b["name"],"subtitle":b["subtitle"],"count":b["count"],"generated":b["id"] in generated,"search_terms":" ".join(r["word"] for r in b["words"])})
    unclassified_count=len(unclassified_batches(u))
    write_json(INDEX_PATH,{"summary":{"total_words":len(all_records),"total_batches":len(batches),"part1_part2_batches":len(batches)-unclassified_count,"unclassified_batches":unclassified_count,"generated_batches":len(ordered)},"batches":metadata})

if __name__=="__main__": main()
